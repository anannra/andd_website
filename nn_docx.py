"""Native Networks document style: shared python-docx helpers.

Arial throughout; ink header rows; horizontal-only hairline tables; indigo
accent bars on callouts; hairline-framed stat tiles. Palette re-tokened from
the v4 whitepaper system to the Native Networks brand.
"""

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ---- tokens (Native Networks, print) ------------------------------------
INK = RGBColor(0x11, 0x1A, 0x30)
BODY = RGBColor(0x3A, 0x44, 0x58)
ACCENT = RGBColor(0x4C, 0x6F, 0xFF)     # indigo
SKY = RGBColor(0x0E, 0x8E, 0xC7)        # sky, darkened for print
GRAY = RGBColor(0x6B, 0x74, 0x88)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xD8, 0xDE, 0xF2)
HAIRLINE = "D5DAE8"
ZEBRA = "F4F6FB"
LIGHT = "EEF1FA"
INKHEX = "111A30"
ACCENTHEX = "4C6FFF"

MARK = "/home/anoop/git/anannra/andd_website/proposals/assets/mark-print.png"
CONTENT_W = 9360  # twips, 6.5in


# ---- helpers ---------------------------------------------------------------
def set_font(run, size=10.5, color=BODY, bold=False, italic=False, spacing=None):
    f = run.font
    f.name = "Arial"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), "Arial")
    f.size = Pt(size)
    f.color.rgb = color
    f.bold = bold
    f.italic = italic
    if spacing:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(spacing))
        rpr.append(sp)


def add_runs(par, text, size=10.5, color=BODY, bold=False, italic=False):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(par.add_run(part[2:-2]), size, INK if color == BODY else color, True, italic)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_font(par.add_run(part[1:-1]), size, color, bold, True)
        else:
            set_font(par.add_run(part), size, color, bold, italic)


def para(doc, text, size=10.5, color=BODY, bold=False, italic=False,
         before=0, after=8, line=1.18, align=None, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if keep_next:
        pf.keep_with_next = True
    if align:
        p.alignment = align
    add_runs(p, text, size, color, bold, italic)
    return p


def eyebrow(doc, text, before=0, after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text.upper()), 8, ACCENT, True, spacing=24)
    return p


def shade(el, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    el.append(shd)


def cell_margins(cell, top=70, bottom=70, left=110, right=110):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, w in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tcpr.append(mar)


def table_borders(tbl, color=HAIRLINE, size=2):
    tblpr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    spec = {"top": "single", "bottom": "single", "insideH": "single",
            "left": "none", "right": "none", "insideV": "none"}
    for side, val in spec.items():
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), val)
        if val == "single":
            b.set(qn("w:sz"), str(size))
            b.set(qn("w:color"), color)
        borders.append(b)
    tblpr.append(borders)


def cell_par(cell, text, size=8.5, color=BODY, bold=False, before=0, after=0, line=1.15):
    p = cell.paragraphs[0] if not cell.paragraphs[0].runs else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    add_runs(p, text, size, color, bold)
    return p


def col_widths(rows, weights=None):
    ncols = len(rows[0])
    if weights is None:
        weights = []
        for c in range(ncols):
            longest = max(len(r[c]) for r in rows if c < len(r))
            weights.append(max(6.0, min(60.0, longest)) ** 0.5)
    total = sum(weights)
    return [int(CONTENT_W * w / total) for w in weights]


def table(doc, rows, weights=None, after=10, size=8.5):
    widths = col_widths(rows, weights)
    tbl = doc.add_table(rows=0, cols=len(rows[0]))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    # keep the grid in step with the cell widths so LibreOffice and Word agree
    for g, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        g.set(qn("w:w"), str(w))
    table_borders(tbl)
    for i, row in enumerate(rows):
        tr = tbl.add_row()
        trpr = tr._tr.get_or_add_trPr()
        if i == 0:
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            trpr.append(th)
        for j, cell_text in enumerate(row):
            cell = tr.cells[j]
            cell.width = Twips(widths[j])
            cell_margins(cell)
            tcpr = cell._tc.get_or_add_tcPr()
            if i == 0:
                shade(tcpr, INKHEX)
                cell_par(cell, cell_text, size, WHITE, True)
            else:
                if i % 2 == 0:
                    shade(tcpr, ZEBRA)
                cell_par(cell, cell_text, size, BODY)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(max(after - 8, 0))
    set_font(sp.add_run(" "), 4, BODY)
    return tbl


def callout(doc, title, body_lines, dark=True, hang=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Twips(CONTENT_W)
    cell_margins(cell, top=140, bottom=140, left=180, right=180)
    tcpr = cell._tc.get_or_add_tcPr()
    shade(tcpr, INKHEX if dark else LIGHT)
    tblpr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "none")
        borders.append(b)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), ACCENTHEX)
    borders.append(left)
    tblpr.append(borders)
    if title:
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(4)
        set_font(p.add_run(title), 11, WHITE if dark else INK, True)
    first = not title
    for line in body_lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.22
        if hang:
            p.paragraph_format.left_indent = Inches(hang)
            p.paragraph_format.first_line_indent = Inches(-hang)
        add_runs(p, line, size=10, color=SOFT if dark else BODY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def stat_tiles(doc, tiles):
    tbl = doc.add_table(rows=1, cols=len(tiles))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tblpr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("left", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "none")
        borders.append(b)
    for side in ("top", "bottom"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), HAIRLINE)
        borders.append(b)
    tblpr.append(borders)
    w = CONTENT_W // len(tiles)
    for j, (num, label) in enumerate(tiles):
        cell = tbl.rows[0].cells[j]
        cell.width = Twips(w)
        cell_margins(cell, top=130, bottom=130, left=40, right=40)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        set_font(p.add_run(num), 15, INK, True)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        set_font(p2.add_run(label.upper()), 7, GRAY, spacing=16)
    return tbl


def hrule(doc, color=ACCENTHEX, size=12, before=2, after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)
    return p


def field(par, instr):
    for tag, attr in (("w:fldChar", {"w:fldCharType": "begin"}),
                      ("w:instrText", None),
                      ("w:fldChar", {"w:fldCharType": "end"})):
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi"):
            fonts.set(qn(a), "Arial")
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "15")
        col = OxmlElement("w:color"); col.set(qn("w:val"), "6B7488")
        rpr.append(fonts); rpr.append(sz); rpr.append(col)
        r.append(rpr)
        el = OxmlElement(tag)
        if attr:
            for k, v in attr.items():
                el.set(qn(k), v)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        r.append(el)
        par._p.append(r)


def page_break(doc):
    brk = doc.add_paragraph()
    brk.paragraph_format.space_after = Pt(0)
    run = brk.add_run()
    run._element.append(OxmlElement("w:br"))
    run._element[-1].set(qn("w:type"), "page")


def h1(doc, num, title, new_page=False):
    if new_page:
        page_break(doc)
    eyebrow(doc, f"{num:02d} · {title}", before=0 if new_page else 26)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(f"{num}. {title}"), 15, INK, True)
    hrule(doc).paragraph_format.keep_with_next = True


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), 11.5, ACCENT, True)


def lead(doc, text):
    para(doc, text, size=11.5, color=INK, bold=len(text) <= 220, line=1.32, after=10, keep_next=True)


def bullets(doc, items, size=10.5):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.18
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.24)
        set_font(p.add_run("•"), size, ACCENT, True)
        add_runs(p, "\t" + item, size, BODY)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.24))


def numbered(doc, items):
    for n, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.18
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        set_font(p.add_run(f"{n:02d}"), 9, ACCENT, True)
        add_runs(p, "   " + item, 10.5, BODY)


def note(doc, text):
    para(doc, text, size=8.5, color=GRAY, italic=True, before=2, after=10)


